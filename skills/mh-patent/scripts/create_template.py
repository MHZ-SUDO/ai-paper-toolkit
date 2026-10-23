#!/usr/bin/env python3
"""Create a sanitized, editable patent working template from bundled resources.

Requires Python 3.10+ and python-docx. Never reads a user's existing invention.
Existing output files are protected unless --overwrite is explicitly supplied.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Twips
except ImportError as exc:
    raise SystemExit("Missing dependency: install python-docx in your selected Python environment.") from exc


RESOURCE_DIR = Path(__file__).resolve().parents[1] / "assets"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def xml_child(parent, tag, **attrs):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    for name, value in attrs.items():
        element.set(qn(name), str(value))
    return element


def font_properties(rpr, latin, chinese):
    fonts = xml_child(rpr, "w:rFonts")
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + name), None)
    for name in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + name), latin)
    fonts.set(qn("w:eastAsia"), chinese)


def setup_styles(doc, profile):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    font_properties(normal.element.get_or_add_rPr(), "Times New Roman", "宋体")
    for entry in profile["styles"]:
        name = entry["role"]
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = entry["font_latin"]
        style.font.size = Pt(entry["font_size_pt"])
        style.font.bold = entry["bold"]
        style.font.color.rgb = RGBColor.from_string("000000")
        rp = style.element.get_or_add_rPr()
        font_properties(rp, entry["font_latin"], entry["font_zh"])
        xml_child(rp, "w:szCs", **{"w:val": round(entry["font_size_pt"] * 2)})
        xml_child(rp, "w:bCs", **{"w:val": "1" if entry["bold"] else "0"})
        pf = style.paragraph_format
        pf.alignment = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }[entry["alignment"]]
        pf.first_line_indent = Pt(entry["first_line_pt"])
        pf.left_indent = Pt(0)
        pf.right_indent = Pt(0)
        pf.space_before = Pt(entry["before_pt"])
        pf.space_after = Pt(entry["after_pt"])
        if entry["line_rule"] == "auto":
            pf.line_spacing = entry["line_value_pt_or_multiplier"]
        else:
            pf.line_spacing = Pt(entry["line_value_pt_or_multiplier"])
            pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.keep_with_next = entry["keep_next"]
        pf.keep_together = entry["keep_lines"]
        pf.widow_control = True
        xml_child(style.element.get_or_add_pPr(), "w:snapToGrid", **{"w:val": "0"})


def configure_document(doc, profile):
    paper = profile["paper"]
    section = doc.sections[0]
    for prop, key in [
        ("page_width", "width_twips"), ("page_height", "height_twips"),
        ("top_margin", "top_twips"), ("bottom_margin", "bottom_twips"),
        ("left_margin", "left_twips"), ("right_margin", "right_twips"),
        ("header_distance", "header_twips"), ("footer_distance", "footer_twips"),
    ]:
        setattr(section, prop, Twips(paper[key]))
    setup_styles(doc, profile)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    footer.paragraph_format.line_spacing = Pt(profile["footer"]["line_pt"])
    footer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    font_properties(rpr, profile["footer"]["font_latin"], profile["footer"]["font_zh"])
    xml_child(rpr, "w:sz", **{"w:val": profile["footer"]["size_pt"] * 2})
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)
    xml_child(doc.settings.element, "w:updateFields", **{"w:val": "true"})
    mp = xml_child(doc.settings.element, "m:mathPr")
    xml_child(mp, "m:mathFont", **{"m:val": profile["formula"]["font"]})
    cp = doc.core_properties
    cp.title = "MH-Patent 通用可编辑工作模板"
    cp.subject = "待填模板：使用前以实际技术方案替换全部占位提示"
    cp.author = "MH-Patent"
    cp.last_modified_by = "MH-Patent"
    cp.keywords = "patent,template,editable"
    cp.comments = "Generic template. Contains no actual invention, drawing, or identity."
    cp.category = "Template"
    cp.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    cp.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)


def create_document(profile, boilerplate):
    doc = Document()
    configure_document(doc, profile)

    def p(text, style="专利正文", page_break=False):
        paragraph = doc.add_paragraph(text, style)
        if page_break:
            paragraph.paragraph_format.page_break_before = True
        return paragraph

    p("【发明名称】", "专利名称")
    p("【完成人信息：仅在内部交底稿中按需填写或删除】", "完成人信息")
    p("说 明 书 摘 要", "专利分部标题")
    p("【摘要待填：说明所属技术领域、要解决的问题、核心技术方案及其作用。仅使用已有技术资料能够支持的内容，完成后删除本提示。】")

    p("摘 要 附 图", "专利分部标题", True)
    image_slot = p("【在此插入能独立表达核心方案的无标号摘要附图】", "专利附图")
    image_slot.paragraph_format.keep_with_next = False

    p("权 利 要 求 书", "专利分部标题", True)
    p("1. 【独立权利要求待填：按实际申请类型写明主题名称、必要技术特征及其关系。结构类写明必要构件与配合关系；控制方法类写明处理步骤、输入输出与执行关系；结合类写明结构与控制的协同关系。】")
    p("2. 【从属权利要求待填：明确引用前项权利要求，并补充已有材料支持的进一步技术特征；按需增减条数。】")

    p("说 明 书", "专利分部标题", True)
    p("【发明名称】", "专利名称")
    p("技术领域", "专利小标题")
    p("【说明实际所属技术领域及具体对象。】")
    p("背景技术", "专利小标题")
    p("【说明与本方案直接相关的现有方案、具体应用条件及待解决的问题；有出处的事实保留可追溯来源。】")
    p("发明内容", "专利小标题")
    p("【依次说明要解决的问题、核心技术方案和由这些特征产生的作用。围绕一个明确主线展开；技术关系应与权利要求及后文实施方式一致。】")
    p("附图说明", "专利小标题")
    p(boilerplate["figure_intro"])
    p("图【编号】为【该附图所表达的结构、信号关系或工作状态】示意图。")
    p("附图标记说明：【编号】—【部件名称】；【编号】—【部件名称】。")
    p("具体实施方式", "专利小标题")
    for entry in boilerplate["opening"]:
        p(entry["text"])
    p("【实施方式待填：从使用场景和初始状态开始，按触发条件、动作或处理过程、状态变化、输出及返回过程连续叙述，并引用相应附图。将理解动作、载荷或信号关系所必需的安装与连接事实嵌入过程，不重复罗列整套结构。】")
    p("【控制类方案应交代测量输入、状态或误差形成、控制计算、执行器作用及反馈；结构控制结合类还应说明控制输出如何引起结构动作，以及实际结构状态如何参与反馈。无相应内容时删除本段提示。】")
    p("【确有必要时在此插入原生公式；无必要公式时删除本段】", "专利公式")
    p("【公式使用时定义符号、方向、单位及适用条件，说明公式与具体工作过程的对应关系；没有公式时删除本段提示。】")
    p(boilerplate["closing"])

    p("说 明 书 附 图", "专利分部标题", True)
    p("【主图图像位置：插入具有必要部件标号的说明书附图】", "专利附图")
    p("图【编号】", "专利图题")
    for letter in ("a", "b", "c"):
        p(f"【子图图像位置：按需插入独立图像；不用子图时删除】", "专利附图")
        p(f"（{letter}）【状态或含义待填】", "专利子图题")
    p("图【编号】", "专利图题")
    return doc


def validate_payload(payload):
    with ZipFile(BytesIO(payload)) as z:
        if z.testzip() is not None:
            raise ValueError("DOCX ZIP integrity check failed")
        document = ET.fromstring(z.read("word/document.xml"))
        texts = [node.text or "" for node in document.findall(".//w:t", NS)]
        joined = "\n".join(texts)
        for heading in ("技术领域", "背景技术", "发明内容", "附图说明", "具体实施方式"):
            if texts.count(heading) != 1:
                raise ValueError(f"Expected exactly one section: {heading}")
        if texts.count("说 明 书 附 图") != 1:
            raise ValueError("Drawing-section heading must occur once")
        if "[0001]" in joined:
            raise ValueError("Unexpected paragraph numbering")
        if document.findall(".//w:drawing", NS):
            raise ValueError("A generic template must not contain invention drawings")
        for name in z.namelist():
            if name.endswith(".rels") and b'TargetMode="External"' in z.read(name):
                raise ValueError("Unexpected external relationship")
    return {"paragraph_count": len(document.findall(".//w:p", NS)), "images": 0, "placeholders_present": "【" in joined}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="New .docx path")
    parser.add_argument("--overwrite", action="store_true", help="Explicitly allow replacing the output file")
    args = parser.parse_args()
    output = args.output.expanduser()
    if output.suffix.lower() != ".docx":
        parser.error("--output must end in .docx")
    if output.exists() and not args.overwrite:
        parser.error("Output already exists; choose a new path or explicitly pass --overwrite")
    profile = json.loads((RESOURCE_DIR / "format-profile.json").read_text(encoding="utf-8"))
    boilerplate = json.loads((RESOURCE_DIR / "boilerplate.json").read_text(encoding="utf-8"))
    doc = create_document(profile, boilerplate)
    buffer = BytesIO()
    doc.save(buffer)
    payload = buffer.getvalue()
    audit = validate_payload(payload)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("wb" if args.overwrite else "xb") as handle:
        handle.write(payload)
    # Read the actual output back instead of treating a successful save as proof.
    validate_payload(output.read_bytes())
    print(json.dumps({"output": str(output.resolve()), "validation": audit}, ensure_ascii=False))


if __name__ == "__main__":
    main()
